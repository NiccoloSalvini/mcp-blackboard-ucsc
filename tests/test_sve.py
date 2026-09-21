"""SVE (Cattolica exam registration) Excel template: fill the Voto Finale
column from Blackboard scores, keep everything else untouched.

Names and student numbers here are invented. Never put a real roster in a test:
this repository is public.
"""
import asyncio

import openpyxl
import pytest

from bb_mcp import client, server, sve


def run(coro):
    return asyncio.run(coro)


def make_template(path, students):
    """A copy of the shape SVE exports: sheet1 with 9 meta rows, headers on row 11,
    data from row 17; sheet2 with the column map SVE itself reads on import."""
    wb = openpyxl.Workbook()
    s1 = wb.active
    s1.title = "Registrazione"
    for i, (k, v) in enumerate([("Facoltà", "ECONOMIA"), ("Sede", "Roma"), ("Insegnamento", "STATISTICS AND BIG DATA"),
                                ("Codice Insegnamento", "SSI428"), ("Sessione", "Autunnale"), ("Anno Accademico", "2025"),
                                ("Appello", "24/08-12/09"), ("Data Appello", "10/09/2026 09:00"),
                                ("Codice Appello", "288551")], start=1):
        s1.cell(i, 1, k); s1.cell(i, 2, v)
    headers = ["Codice Studente", "Codice Insegnamento", "Insegnamento", "Codice indirizzo", "Codice sotto indirizzo",
               "Descrizione indirizzo", "Descrizione sotto indirizzo", "CFU", "A.A. Freq.", "L / IC / EA", "Matricola",
               "Cognome Nome", "Voto Finale", "Matricola Docente o Collaboratore", "Data Svolgimento Esame",
               "Voto 2", "Voto 3", "Voto 4", "Voto 5", "Note", "Email Studente"]
    for c, h in enumerate(headers, start=1):
        s1.cell(11, c, h)
    for r, (sid, matricola, name) in enumerate(students, start=17):
        s1.cell(r, 1, sid); s1.cell(r, 2, "SSI428"); s1.cell(r, 3, "STATISTICS AND BIG DATA")
        s1.cell(r, 8, "08.00"); s1.cell(r, 11, matricola); s1.cell(r, 12, name); s1.cell(r, 21, "x@icatt.it")
    s2 = wb.create_sheet("Codici di controllo")
    # rows 1-9: metadata map — B is the row, C the column, both 0-based, of sheet 1
    for r, (key, row0) in enumerate([("COD_INSEGNAMENTO", 3), ("DES_INSEGNAMENTO", 2), ("FRAZIONE", 6),
                                     ("DES_SESSIONE", 4), ("DATA_APPELLO", 7), ("SEDE", 1), ("FACOLTA", 0),
                                     ("ID_APPELLO", 8), ("ANNO_ACCADEMICO", 5)], start=1):
        s2.cell(r, 1, key); s2.cell(r, 2, row0); s2.cell(r, 3, 1)
    s2["A11"] = "LIST_HEADER"; s2["B11"] = 10
    for r, (key, idx) in enumerate([("CODICE_STUDENTE", 0), ("MATRICOLA", 10), ("COGNOME_NOME", 11), ("VOTO", 12),
                                    ("DATA_SVOLGIMENTO_ESAME", 14), ("NOTE", 19)], start=12):
        s2.cell(r, 1, key); s2.cell(r, 3, idx)
    s2["A33"] = "LIST_CONTENT"; s2["B33"] = 16
    s2["A36"] = "UNLOCK_START"; s2["C36"] = 14
    s2["A37"] = "UNLOCK_END"; s2["C37"] = 19
    wb.save(path)


STUDENTS = [("9000001", "1111111", "ROSSI MARIA"), ("9000002", "2222222", "BIANCHI LUCA"),
            ("9000003", "3333333", "VERDI SARA")]


def test_layout_is_read_from_the_config_sheet(tmp_path):
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)
    lay = sve.read_layout(openpyxl.load_workbook(p))
    assert lay.header_row == 11 and lay.first_row == 17
    assert lay.col("MATRICOLA") == 11 and lay.col("VOTO") == 13 and lay.col("NOTE") == 20
    assert lay.appello == "288551" and lay.insegnamento == "SSI428"
    # the config sheet also maps metadata and an unlocked range: neither is a data column
    assert "SEDE" not in lay.columns and "UNLOCK_START" not in lay.columns
    assert lay.meta["FACOLTA"] == "ECONOMIA" and lay.meta["DES_SESSIONE"] == "Autunnale"
    with pytest.raises(client.BlackboardError, match="no SEDE column"):
        lay.col("SEDE")


def test_matricola_is_matched_even_when_excel_stored_it_as_a_number(tmp_path):
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)
    wb = openpyxl.load_workbook(p); wb.worksheets[0]["K17"] = 1111111; wb.save(p)   # int, not text
    rep = sve.fill(p, {"1111111": 30}, tmp_path / "o.xlsx")
    assert rep["filled"] == 1 and rep["graded_but_not_enrolled"] == []


def test_grade_to_esito_mapping():
    assert sve.esito(30) == "30"
    assert sve.esito(26.25) == "26"
    assert sve.esito(27.5) == "28"
    assert sve.esito(31) == "30 e lode"
    assert sve.esito(17.4) == "non sufficiente"
    assert sve.esito(17.5) == "18"          # pass/fail is decided on the rounded mark
    assert sve.esito(17.4, not_passed="respinto") == "respinto"
    assert sve.esito(None) is None
    with pytest.raises(client.BlackboardError, match="outside"):
        sve.esito(-1)


def test_fill_writes_voto_finale_and_reports_mismatches(tmp_path):
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)
    out = tmp_path / "out.xlsx"
    rep = sve.fill(p, {"1111111": 28.4, "2222222": 12.0, "9999999": 30}, out, exam_date="10/09/2026")
    wb = openpyxl.load_workbook(out); s1 = wb.worksheets[0]
    assert s1["M17"].value == "28" and s1["M18"].value == "non sufficiente" and s1["M19"].value is None
    assert s1["O17"].value == "10/09/2026" and s1["O19"].value is None
    assert s1["L17"].value == "ROSSI MARIA" and s1["A17"].value == "9000001"        # untouched
    assert wb.worksheets[1]["A11"].value == "LIST_HEADER"                             # config sheet kept
    assert rep == {"path": str(out), "appello": "288551", "insegnamento": "SSI428", "rows": 3, "filled": 2,
                   "esiti": {"28": 1, "non sufficiente": 1},
                   "no_grade_in_blackboard": [{"matricola": "3333333", "name": "VERDI SARA"}],
                   "graded_but_not_enrolled": ["9999999"],
                   "next": "SVE > Registra voti > Caricamento voti Excel > Importa esiti da XLS, check the Errore "
                           "column, then Inoltro alla firma."}


def test_fill_refuses_to_overwrite_an_existing_esito_unless_asked(tmp_path):
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)
    wb = openpyxl.load_workbook(p); wb.worksheets[0]["M17"] = "30"; wb.save(p)
    with pytest.raises(client.BlackboardError, match="already has"):
        sve.fill(p, {"1111111": 20}, tmp_path / "o.xlsx")
    rep = sve.fill(p, {"1111111": 20}, tmp_path / "o.xlsx", overwrite=True)
    assert rep["filled"] == 1


def test_tool_reads_column_and_fills(monkeypatch, tmp_path):
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)
    calls = []

    async def fake_grades(course_id, column_id):
        calls.append((course_id, column_id))
        return {"column": {"id": column_id, "name": "Full exam"}, "rows": [
            {"userId": "_9_1", "name": "Rossi Maria", "studentId": "1111111", "score": 30.0},
            {"userId": "_8_1", "name": "Verdi Sara", "studentId": "3333333", "score": None}]}

    async def no_attempts(course_id, column_id):
        return []

    monkeypatch.setattr(server, "bb_list_grades", fake_grades)
    monkeypatch.setattr(server, "bb_list_attempts", no_attempts)
    rep = run(server.bb_sve_fill_excel("_1_1", "_g_1", str(p)))
    assert calls == [("_1_1", "_g_1")]
    assert rep["filled"] == 1 and rep["path"].endswith("t-compilato.xlsx")
    assert [r["matricola"] for r in rep["no_grade_in_blackboard"]] == ["2222222", "3333333"]


def test_tool_flags_submissions_still_waiting_to_be_marked(monkeypatch, tmp_path):
    """A student enrolled in the appello who handed in but has no mark yet must be
    shouted about: verbalising that row blank is the costly mistake here."""
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)

    async def fake_grades(course_id, column_id):
        return {"column": {"id": column_id, "name": "Full exam"}, "rows": [
            {"userId": "_9_1", "name": "Rossi Maria", "studentId": "1111111", "score": 30.0},
            {"userId": "_8_1", "name": "Bianchi Luca", "studentId": "2222222", "score": None},
            {"userId": "_7_1", "name": "Verdi Sara", "studentId": "3333333", "score": None}]}

    async def fake_attempts(course_id, column_id):
        return [{"attemptId": "_a_1", "userId": "_8_1", "status": "NeedsGrading",
                 "attemptDate": "2026-09-16T11:26:19.669Z"}]

    monkeypatch.setattr(server, "bb_list_grades", fake_grades)
    monkeypatch.setattr(server, "bb_list_attempts", fake_attempts)
    rep = run(server.bb_sve_fill_excel("_1_1", "_g_1", str(p)))
    assert rep["filled"] == 1
    assert rep["submitted_not_yet_marked"] == [
        {"matricola": "2222222", "name": "Bianchi Luca", "attemptId": "_a_1",
         "attemptDate": "2026-09-16T11:26:19.669Z"}]
    assert rep["warning"].startswith("1 submission")
    # the student who simply never sat the exam is only in the quiet list
    assert [r["matricola"] for r in rep["no_grade_in_blackboard"]] == ["2222222", "3333333"]


def test_tool_has_no_warning_when_nothing_is_pending(monkeypatch, tmp_path):
    p = tmp_path / "t.xlsx"; make_template(p, STUDENTS)

    async def fake_grades(course_id, column_id):
        return {"column": {"id": column_id, "name": "Full exam"}, "rows": [
            {"userId": "_9_1", "name": "Rossi Maria", "studentId": "1111111", "score": 30.0}]}

    async def fake_attempts(course_id, column_id):
        return []

    monkeypatch.setattr(server, "bb_list_grades", fake_grades)
    monkeypatch.setattr(server, "bb_list_attempts", fake_attempts)
    rep = run(server.bb_sve_fill_excel("_1_1", "_g_1", str(p)))
    assert rep["submitted_not_yet_marked"] == [] and "warning" not in rep


# --- office-hours booking via self-enrolment groups -------------------------------

def test_group_set_can_be_self_enrol_with_limit(monkeypatch):
    monkeypatch.setattr(client, "ALLOW_WRITES", True)
    routes = {
        ("GET", "/v1/courses/_1_1/users"): {"results": []},
        ("POST", "/v2/courses/_1_1/groups/sets"): {"id": "_set_1"},
        ("POST", "/v2/courses/_1_1/groups/sets/_set_1/groups"): {"id": "_g_x"},
    }
    calls = []

    async def req(method, path, *, json=None, params=None, **kw):
        calls.append((method, path, json)); return routes[(method, path)]

    monkeypatch.setattr(client, "_request", req)
    out = run(server.bb_create_group_set("_1_1", "Visione compiti 10/09", [{"name": "14:00"}, {"name": "14:20"}],
                                         available=True, self_enroll=True, limit=3))
    assert calls[1][2]["enrollment"] == {"type": "SelfEnrollment", "limit": 3,
                                         "signupSheet": {"name": "Visione compiti 10/09", "showMembers": False}}
    assert calls[1][2]["availability"] == {"available": "SignupOnly"}
    assert calls[2][2]["enrollment"]["type"] == "SelfEnrollment" and calls[2][2]["enrollment"]["limit"] == 3
    assert out["set_id"] == "_set_1"
